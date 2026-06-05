"""utils/helpers pytest 测试"""
import os
import json
import pytest
from utils.helpers import (
    save_result, load_companies_from_file, export_csv,
    export_excel, export_batch_excel, _classify_phone, setup_logger,
)


class TestSaveResult:
    def test_save_to_dir(self, tmp_dir):
        path = save_result("测试公司", [{"phone": "138"}], tmp_dir)
        assert os.path.exists(path)
        with open(path) as f:
            data = json.load(f)
        assert data["company"] == "测试公司"
    
    def test_save_to_file(self, tmp_dir):
        path = save_result("ABC", [{"x": 1}], os.path.join(tmp_dir, "abc.json"))
        assert os.path.exists(path)


class TestLoadCompanies:
    def test_normal(self, tmp_dir):
        f = os.path.join(tmp_dir, "c.txt")
        with open(f, "w") as fh:
            fh.write("A公司\nB公司\n# 注释\n\nC公司\n")
        result = load_companies_from_file(f)
        assert result == ["A公司", "B公司", "C公司"]
    
    def test_not_found(self):
        with pytest.raises(FileNotFoundError):
            load_companies_from_file("/nonexistent/path.txt")


class TestClassifyPhone:
    def test_mobile(self):
        assert _classify_phone("13800138000") == "手机"
    
    def test_landline(self):
        assert _classify_phone("053112345") == "固话"
    
    def test_400(self):
        assert _classify_phone("4001234567") == "服务热线"
    
    def test_800(self):
        assert _classify_phone("8001234567") == "服务热线"


class TestExport:
    def test_csv(self, tmp_dir, sample_sources):
        path = export_csv("测试", ["13800138000"], sample_sources, os.path.join(tmp_dir, "t.csv"))
        assert os.path.exists(path)
    
    def test_excel(self, tmp_dir, sample_sources):
        path = export_excel("测试", ["13800138000"], sample_sources, os.path.join(tmp_dir, "t.xlsx"))
        assert os.path.exists(path)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert wb.active.title == "电话列表"
    
    def test_batch_excel(self, tmp_dir):
        data = [
            {"company": "A", "phones": ["13800138000"], "sources": [{"phone": "13800138000", "source": "百度地图POI"}]},
            {"company": "B", "phones": [], "sources": []},
        ]
        path = export_batch_excel(data, os.path.join(tmp_dir, "batch.xlsx"))
        assert os.path.exists(path)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert "搜索结果" in wb.sheetnames
        assert "统计汇总" in wb.sheetnames
        assert "未找到" in wb.sheetnames
