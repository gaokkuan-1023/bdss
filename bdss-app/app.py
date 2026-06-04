#!/usr/bin/env python3
"""BDSS 批量获客系统 — 轻量版（无需 Playwright）"""
import sys, os, json, time, threading, uuid, logging
logger = logging.getLogger(__name__)
from pathlib import Path

sys.path.insert(0, os.path.expanduser("~/bdss-op"))

from flask import Flask, request, jsonify, render_template_string, send_file
from openpyxl import load_workbook, Workbook
from datetime import datetime
from skills.lite_search import search_company_phones
from bidding_monitor import scan_with_notify, get_stats, update_status
from industry_api import bp as industry_bp

app = Flask(__name__)
app.register_blueprint(industry_bp)
UPLOAD_DIR = Path("/tmp/bdss_uploads")
UPLOAD_DIR.mkdir(exist_ok=True)
tasks = {}

# ============ Dashboard API ============

@app.route("/api/dashboard")
def api_dashboard():
    """仪表盘数据"""
    from bidding_monitor import get_db, close_db
    conn = get_db()
    total = conn.execute("SELECT COUNT(*) FROM bidding_items").fetchone()[0]
    new_today = conn.execute("SELECT COUNT(*) FROM bidding_items WHERE matched_at > ?", (time.time() - 86400,)).fetchone()[0]
    with_phone = conn.execute("SELECT COUNT(*) FROM bidding_items WHERE phone != ''").fetchone()[0]
    contacted = conn.execute("SELECT COUNT(*) FROM bidding_items WHERE status='contacted'").fetchone()[0]
    close_db()
    return jsonify({"total": total, "new_today": new_today, "with_phone": with_phone, "contacted": contacted})


# ============ 关键词管理 API ============

@app.route("/api/keywords", methods=["GET", "POST"])
def api_keywords():
    from bidding_monitor import CCGP_KEYWORDS, OKCIS_KEYWORDS, WATER_CHEM_KEYWORDS
    if request.method == "POST":
        data = request.get_json(silent=True) or {}
        new_kw = data.get("keywords", [])
        # 持久化到文件
        kw_path = Path.home() / "bdss-op/water_chem_keywords.json"
        with open(kw_path, "w") as f:
            json.dump(new_kw, f, ensure_ascii=False)
        return jsonify(success=True, keywords=new_kw)
    # GET: 返回当前关键词
    return jsonify({
        "ccgp": CCGP_KEYWORDS,
        "okcis": OKCIS_KEYWORDS,
        "all": WATER_CHEM_KEYWORDS,
    })


# ============ 公司详情 API ============

@app.route("/api/company/<path:name>")
def api_company(name):
    from bidding_monitor import get_db, close_db
    conn = get_db()
    rows = conn.execute(
        "SELECT id, title, source, buyer, contact, phone, email, matched_at, status FROM bidding_items WHERE buyer LIKE ? ORDER BY matched_at DESC LIMIT 20",
        (f"%{name}%",),
    ).fetchall()
    return jsonify([{"id": r[0], "title": r[1], "source": r[2], "buyer": r[3],
                     "contact": r[4], "phone": r[5], "email": r[6],
                     "time": r[7], "status": r[8]} for r in rows])

# ============ 外呼系统 API ============

@app.route("/api/outcall/config", methods=["GET", "POST"])
def api_outcall_config():
    """配置阿里云外呼接口"""
    config_path = Path(__file__).parent / "outcall_config.json"
    if request.method == "POST":
        data = request.get_json(silent=True) or {}
        with open(config_path, "w") as f:
            json.dump(data, f, ensure_ascii=False)
        return jsonify(success=True)
    if config_path.exists():
        with open(config_path) as f:
            return jsonify(json.load(f))
    return jsonify({"enabled": False, "api_url": "", "api_key": ""})

@app.route("/api/outcall/push", methods=["POST"])
def api_outcall_push():
    """推送客户名单到阿里云外呼"""
    body = request.get_json(silent=True) or {}
    contacts = body.get("contacts", [])
    if not contacts:
        return jsonify(success=False, error="没有要推送的联系人")
    config_path = Path(__file__).parent / "outcall_config.json"
    config = {"enabled": False, "api_url": "", "api_key": ""}
    if config_path.exists():
        with open(config_path) as f:
            config = json.load(f)
    if not config.get("enabled") or not config.get("api_url"):
        return jsonify(success=False, error="未配置外呼系统，请先在设置中配置")
    # 构造请求
    payload = {
        "contacts": [{"name": c.get("name",""), "phone": c.get("phone","")} for c in contacts if c.get("phone")],
        "total": len(contacts),
    }
    import urllib.request as _req
    try:
        data = json.dumps(payload, ensure_ascii=False).encode()
        req = _req.Request(config["api_url"], data=data,
            headers={"Content-Type": "application/json", "Authorization": f"Bearer {config.get('api_key','')}"})
        resp = _req.urlopen(req, timeout=15)
        result = json.loads(resp.read())
        return jsonify(success=True, result=result, pushed=len(payload["contacts"]))
    except Exception as e:
        return jsonify(success=False, error=f"推送失败: {str(e)[:100]}")

@app.route("/api/outcall/push_bidding", methods=["POST"])
def api_outcall_push_bidding():
    """推送招标线索到外呼"""
    from bidding_monitor import get_stats
    stats = get_stats()
    contacts = []
    for item in stats["recent"]:
        if item.get("phone") and item.get("status") == "new":
            contacts.append({"name": item.get("buyer","") or item.get("title","")[:20], "phone": item["phone"]})
    if not contacts:
        return jsonify(success=False, error="没有未联系且有电话的线索")
    # 构造请求体
    from flask import request as _req_ctx
    with app.test_request_context(json={"contacts": contacts}):
        return api_outcall_push()
    pois = [p for p in pois if not p["name"].endswith(("市","省","区","县"))]
    if not pois:
        alt_keywords = [f"{keyword}公司", f"{keyword}厂", f"{keyword}企业"]
        import os as _os
        ak = _os.environ.get("BAIDU_MAP_AK", "")
        if not ak:
            try:
                from utils.env import get_ak as _get_ak
                ak = _get_ak()
            except:
                pass
        if ak:
            import urllib.request, json as _json, urllib.parse
            for alt_kw in alt_keywords:
                url = f"https://api.map.baidu.com/place/v2/search?query={urllib.parse.quote(alt_kw)}&region={urllib.parse.quote(region)}&output=json&ak={ak}&page_size=5"
                try:
                    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                    resp = urllib.request.urlopen(req, timeout=8)
                    data = _json.loads(resp.read())
                    for poi in data.get("results", []):
                        name = poi.get("name", "")
                        if not name or name.endswith(("市","省","区","县")):
                            continue
                        if not any(p["name"] == name for p in pois):
                            pois.append({"name": name, "address": poi.get("address",""), "phone": poi.get("telephone","") or poi.get("phone",""), "uid": poi.get("uid","")})
                except:
                    pass
    return jsonify(success=True, companies=pois[:30], keyword=keyword, region=region)
INDUSTRY_HTML = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>BDSS 行业拓客</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;background:#0f172a;color:#e2e8f0;padding:24px}
.nav{display:flex;gap:16px;margin-bottom:24px;align-items:center}
.nav a{color:#94a3b8;text-decoration:none;padding:6px 14px;border-radius:6px;font-size:.9rem}
.nav a:hover{color:#e2e8f0;background:rgba(255,255,255,.04)}
.nav a.active{color:#38bdf8;background:rgba(56,189,248,.1)}
h1{font-size:1.3rem;font-weight:700;margin-bottom:4px;background:linear-gradient(135deg,#38bdf8,#a78bfa);-webkit-background-clip:text;-webkit-text-fill-color:transparent}
.sub{color:#94a3b8;font-size:.85rem;margin-bottom:16px}
.card{background:#1e293b;border:1px solid #334155;border-radius:10px;padding:20px;margin-bottom:12px}
.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center}
.row input{flex:1;min-width:150px;padding:8px 12px;background:#0f172a;border:1px solid #334155;border-radius:6px;color:#e2e8f0;font-size:.85rem}
.row button{padding:8px 20px;background:#2563eb;color:#fff;border:none;border-radius:6px;font-weight:600;cursor:pointer;font-size:.85rem}
.row button:disabled{opacity:.4;cursor:wait}
table{width:100%;border-collapse:collapse;font-size:.82rem}
th,td{padding:7px 10px;text-align:left;border-bottom:1px solid #1e293b}
th{color:#94a3b8;font-weight:600;font-size:.75rem}
td{color:#e2e8f0}
.phone{color:#22c55e}
.btn-sm{padding:4px 10px;border-radius:4px;font-size:.72rem;cursor:pointer;border:1px solid #334155;background:transparent;color:#94a3b8}
.btn-sm:hover{border-color:#38bdf8;color:#38bdf8}
#log{font-size:.78rem;color:#94a3b8;background:#0f172a;padding:10px;border-radius:6px;margin-top:8px;max-height:150px;overflow-y:auto}
</style>
</head>
<body>
<div class=nav><a href=/ class="nav-item active">\U0001f4e4 \u641c\u7535\u8bdd</a><a href=/industry class="nav-item">\U0001f3ed \u62d3\u5ba2</a><a href=/monitor class="nav-item">\U0001f4e1 \u62db\u6807</a><span style="flex:1"></span><span style="color:#38bdf8;font-weight:800;font-size:14px">BDSS</span></div>
<h1>\U0001f3ed \u884c\u4e1a\u62d3\u5ba2</h1>
<p class=sub>\u8f93\u5165\u884c\u4e1a\u5173\u952e\u8bcd\uff0c\u81ea\u52a8\u641c\u7d22\u76ee\u6807\u516c\u53f8\u540d\u5355</p>
<div class=card>
<div class=row>
<input id=kwInput placeholder="\u4f8b\u5982\uff1a\u6c34\u5904\u7406\u836f\u5242\u3001\u706b\u529b\u53d1\u7535\u3001\u70ed\u7535" value="\u6c34\u5904\u7406\u836f\u5242">
<select id=regionSelect style="padding:8px;background:#0f172a;border:1px solid #334155;border-radius:6px;color:#e2e8f0">
<option>\u5c71\u4e1c</option><option>\u6cb3\u5317</option><option>\u6c5f\u82cf</option><option>\u6d59\u6c5f</option><option>\u5e7f\u4e1c</option><option>\u6cb3\u5357</option><option>\u5168\u56fd</option>
</select>
<button onclick=search() id=searchBtn>\U0001f50d \u641c\u7d22</button>
</div>
<table id=resultTable style=display:none;margin-top:12px>
<thead><tr><th>\u516c\u53f8\u540d\u79f0</th><th>\u7535\u8bdd</th><th>\u5730\u5740</th><th>\u64cd\u4f5c</th></tr></thead>
<tbody id=resultBody></tbody>
</table>
<div id=log></div>
</div>
<script>
async function search(){
const kw=document.getElementById('kwInput').value.trim();
const region=document.getElementById('regionSelect').value;
if(!kw)return;
document.getElementById('searchBtn').disabled=true;
document.getElementById('searchBtn').textContent='\u23f3 \u641c\u7d22\u4e2d...';
document.getElementById('resultTable').style.display='none';
document.getElementById('log').textContent='\u641c\u7d22\u4e2d...\n';
const r=await fetch('/api/industry_search?keyword='+encodeURIComponent(kw)+'&region='+encodeURIComponent(region));
const d=await r.json();
document.getElementById('log').textContent='\u627e\u5230 '+d.companies.length+' \u5bb6\u516c\u53f8\n';
document.getElementById('resultBody').innerHTML=d.companies.map(c=>'<tr><td>'+c.name+'</td><td class=phone>'+(c.phone||'')+'</td><td>'+(c.address||'')+'</td><td><button class=btn-sm onclick=\\"lookupPhone(\\''+encodeURIComponent(c.name)+'\\')\\">\U0001f50d \u67e5\u7535\u8bdd</button></td></tr>').join('');
document.getElementById('resultTable').style.display='table';
document.getElementById('searchBtn').disabled=false;
document.getElementById('searchBtn').textContent='\U0001f50d \u641c\u7d22';
}
async function pushToOutcall(contacts){
const r=await fetch('/api/outcall/push',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({contacts})});
async function lookupPhone(name){window.location.href='/?company='+encodeURIComponent(name);}
</script>
</body></html>"""


@app.route("/industry")
def industry_page():
    return render_template_string(INDUSTRY_HTML)
INDEX_HTML = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>BDSS 批量获客系统</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;background:#0f172a;color:#e2e8f0;padding:24px}
.container{max-width:860px;margin:0 auto}
.nav{display:flex;gap:16px;margin-bottom:16px;align-items:center}
.nav a{color:#94a3b8;text-decoration:none;padding:6px 14px;border-radius:6px;font-size:.9rem}
.nav a:hover{color:#e2e8f0;background:rgba(255,255,255,.04)}
.nav a.active{color:#38bdf8;background:rgba(56,189,248,.1)}
h1{font-size:1.8rem;font-weight:800;text-align:center;background:linear-gradient(135deg,#38bdf8,#a78bfa);-webkit-background-clip:text;-webkit-text-fill-color:transparent;margin-bottom:4px}
.sub{text-align:center;color:#94a3b8;margin-bottom:24px;font-size:.9rem}
.card{background:#1e293b;border:1px solid #334155;border-radius:12px;padding:24px;margin-bottom:16px}
.card h2{font-size:1rem;margin-bottom:12px;display:flex;align-items:center;gap:8px}
.card h2::before{content:'';display:inline-block;width:3px;height:16px;background:#38bdf8;border-radius:2px}
.upload-zone{border:2px dashed #334155;border-radius:12px;padding:40px;text-align:center;cursor:pointer;transition:.2s}
.upload-zone:hover{border-color:#38bdf8;background:rgba(56,189,248,.03)}
.upload-zone .icon{font-size:2.5rem;margin-bottom:8px}
.upload-zone p{color:#94a3b8;font-size:.9rem}
.upload-zone .hint{color:#475569;font-size:.78rem;margin-top:6px}
input[type="file"]{display:none}
.btn{width:100%;padding:12px;border-radius:8px;font-size:.95rem;font-weight:600;cursor:pointer;border:none;transition:.2s;display:inline-block;text-decoration:none;text-align:center}
.btn-primary{background:#2563eb;color:#fff}
.btn-primary:hover{background:#1d4ed8}
.btn-primary:disabled{opacity:.4;cursor:not-allowed}
.btn-success{background:#059669;color:#fff}
.btn-success:hover{background:#047857}
.stats{display:grid;grid-template-columns:repeat(4,1fr);gap:8px;margin:12px 0}
.stat{text-align:center;padding:12px;background:#0f172a;border-radius:8px}
.stat .num{font-size:1.3rem;font-weight:700;color:#38bdf8}
.stat .label{font-size:.72rem;color:#64748b;margin-top:2px}
.log-area{background:#0f172a;border-radius:8px;padding:12px;height:250px;overflow-y:auto;font-family:monospace;font-size:.78rem;line-height:1.6;margin-top:12px}
.log-area .found{color:#22c55e}
.log-area .empty{color:#f59e0b}
.log-area .info{color:#94a3b8}
.progress-bar{background:#0f172a;border-radius:8px;height:8px;overflow:hidden;margin-bottom:8px}
.progress-bar .fill{height:100%;background:linear-gradient(90deg,#38bdf8,#a78bfa);border-radius:8px;width:0%;transition:width .3s}
.table-wrap{overflow-x:auto;font-size:.82rem;margin-top:12px}
table{width:100%;border-collapse:collapse}
th,td{padding:8px 10px;text-align:left;border-bottom:1px solid #1e293b;white-space:nowrap;max-width:200px;overflow:hidden;text-overflow:ellipsis}
th{color:#94a3b8;font-weight:600;font-size:.78rem}
td{font-size:.82rem}
@media(max-width:640px){.stats{grid-template-columns:repeat(2,1fr)}}
</style>
<div class=nav><a href=/ class="nav-item active">\U0001f4e4 \u641c\u7535\u8bdd</a><a href=/industry class="nav-item">\U0001f3ed \u62d3\u5ba2</a><a href=/monitor class="nav-item">\U0001f4e1 \u62db\u6807</a><span style="flex:1"></span><span style="color:#38bdf8;font-weight:800;font-size:14px">BDSS</span></div>
<h1>📞 BDSS 批量获客系统</h1>
<p class="sub">上传公司名单 → 自动百度地图+全网搜电话 → 下载结果</p>
<div class="card">
<h2>📤 第一步：上传公司名单</h2>
<div class="upload-zone" onclick="document.getElementById('fileInput').click()">
<div class="icon">📄</div>
<p>点击上传 Excel 文件</p>
<p class="hint">支持 .xlsx / .xls，需包含「公司名称」列</p>
</div>
<input type="file" id="fileInput" accept=".xlsx,.xls" onchange="handleFile(this)" style="display:none">
<div id="preview" style="display:none;margin-top:16px">
<p id="companyCount" style="color:#94a3b8;margin-bottom:8px"></p>
<div class="table-wrap" id="previewTable"></div>
<button class="btn btn-primary" id="startBtn" onclick="startSearch()" style="margin-top:16px">🚀 开始搜索</button>
</div>
</div>
<div class="card" id="progressCard" style="display:none">
<h2>⏳ 搜索进度</h2>
<div class="stats" id="stats"></div>
<div class="progress-bar"><div class="fill" id="progressFill"></div></div>
<div class="log-area" id="logArea"></div>
<a class="btn btn-success" id="downloadBtn" style="display:none;margin-top:12px" onclick="downloadResult()">📥 下载结果 Excel</a>
</div>
<div class="card" style="cursor:default;font-size:.82rem;color:#64748b">
<p><strong>数据来源：</strong>百度地图 POI · 必应搜索</p>
<p><strong>无需安装浏览器，配置文件即用。</strong></p>
</div>
</div>
<script>
let taskId = null, pollTimer = null;
function handleFile(input) {
const file = input.files[0]; if (!file) return;
const fd = new FormData(); fd.append('file', file);
fetch('/upload', {method:'POST',body:fd}).then(r=>r.json()).then(d=>{
if(!d.success){alert(d.error);return}
document.getElementById('preview').style.display='block';
document.getElementById('companyCount').textContent='共 '+d.count+' 家公司';
let h='<table><thead><tr>'+d.columns.map(c=>'<th>'+c+'</th>').join('')+'<th>操作</th></tr></thead><tbody>';
d.rows.forEach(r=>{h+='<tr>'+r.map(c=>'<td>'+(c||'')+'</td>').join('')+'<td>待搜索</td></tr>'});
h+='</tbody></table>';
document.getElementById('previewTable').innerHTML=h;
taskId=d.task_id;
}).catch(e=>alert('失败: '+e.message));
}
function startSearch(){
document.getElementById('startBtn').disabled=true;document.getElementById('startBtn').textContent='⏳ 搜索中...';
document.getElementById('progressCard').style.display='block';
fetch('/start/'+taskId,{method:'POST'});
pollTimer=setInterval(()=>{
fetch('/progress/'+taskId).then(r=>r.json()).then(d=>{
document.getElementById('progressFill').style.width=d.percent+'%';
document.getElementById('stats').innerHTML=
'<div class=stat><div class=num>'+d.total+'</div><div class=label>总公司</div></div>'+
'<div class=stat><div class=num>'+d.done+'</div><div class=label>完成</div></div>'+
'<div class=stat><div class=num>'+d.phones+'</div><div class=label>电话数</div></div>'+
'<div class=stat><div class=num>'+d.failed+'</div><div class=label>未找到</div></div>';
const log=document.getElementById('logArea');
d.logs.slice(-20).forEach(l=>{if(!log.querySelector('[data-i="'+l.i+'"]')){const e=document.createElement('div');e.setAttribute('data-i',l.i);e.className=l.t;e.textContent=l.m;log.appendChild(e);log.scrollTop=log.scrollHeight}});
if(d.done>=d.total){clearInterval(pollTimer);document.getElementById('startBtn').style.display='none';document.getElementById('downloadBtn').style.display='block'}
});
},1000);
}
function downloadResult(){window.location.href='/download/'+taskId}
</script>
</body>
</html>"""

MONITOR_HTML = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>BDSS 招标监控</title>
<style>
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;background:#0f172a;color:#e2e8f0;padding:24px}
.nav{display:flex;gap:16px;margin-bottom:24px;align-items:center}
.nav a{color:#94a3b8;text-decoration:none;padding:6px 14px;border-radius:6px;font-size:.9rem}
.nav a:hover{color:#e2e8f0;background:rgba(255,255,255,.04)}
.nav a.active{color:#38bdf8;background:rgba(56,189,248,.1)}
h1{font-size:1.5rem;font-weight:800;background:linear-gradient(135deg,#38bdf8,#a78bfa);-webkit-background-clip:text;-webkit-text-fill-color:transparent}
.card{background:#1e293b;border:1px solid #334155;border-radius:10px;padding:20px;margin-bottom:16px}
.card h2{font-size:1rem;margin-bottom:12px;color:#38bdf8}
.stats{display:grid;grid-template-columns:repeat(4,1fr);gap:8px;margin-bottom:16px}
.stat{text-align:center;padding:14px;background:#0f172a;border-radius:8px}
.stat .num{font-size:1.5rem;font-weight:700;color:#38bdf8}
.stat .label{font-size:.75rem;color:#64748b}
.btn{padding:8px 20px;border-radius:6px;border:none;font-weight:600;cursor:pointer;font-size:.85rem}
.btn-primary{background:#2563eb;color:#fff}
.btn-primary:hover{background:#1d4ed8}
.btn-primary:disabled{opacity:.4;cursor:wait}.btn-ghost{background:transparent;border:1px solid #334155;color:#94a3b8;padding:4px 10px;border-radius:4px;font-size:.76rem;cursor:pointer}.btn-ghost:hover{border-color:#38bdf8;color:#38bdf8}
.item{border-bottom:1px solid #1e293b;padding:10px 0;display:flex;justify-content:space-between;align-items:flex-start;gap:12px}
.item:last-child{border:none}
.item .t{flex:1}
.item .t .title{font-size:.88rem;margin-bottom:2px}
.item .t .meta{font-size:.73rem;color:#64748b}
.tag{font-size:.7rem;padding:2px 7px;border-radius:4px}
.tag-new{background:rgba(56,189,248,.12);color:#38bdf8}
.tag-done{background:rgba(34,197,94,.12);color:#22c55e}
#logArea{font-size:.78rem;color:#94a3b8;margin-top:10px}
.keyword-tag{display:inline-block;padding:4px 10px;border-radius:14px;font-size:.78rem;background:rgba(56,189,248,.08);color:#38bdf8;margin:3px}
</style>
</head>
<body>
<div class=nav><a href=/ class="nav-item">\U0001f4e4 \u641c\u7535\u8bdd</a><a href=/industry class="nav-item">\U0001f3ed \u62d3\u5ba2</a><a href=/monitor class="nav-item active">\U0001f4e1 \u62db\u6807</a><span style="flex:1"></span><span style="color:#38bdf8;font-weight:800;font-size:14px">BDSS</span></div>
<div class=stats id=statsArea></div>
<div id=keywordArea></div>
</div>
<div class=card><h2>📋 最新线索</h2><div id=itemList><p style=color:#64748b>加载中...</p></div>
<div style=display:flex;gap:8px;margin-top:12px;flex-wrap:wrap><a href=/api/monitor/export class=btn btn-success style=font-size:.8rem;padding:6px 14px;width:auto>📥 导出 Excel</a><button class="btn btn-ghost" onclick=showKeywords() style=font-size:.8rem;padding:6px 14px>⚙️ 关键词</button><span id=scanStatus style=color:#64748b;font-size:.78rem;align-self:center></span></div></div>
<div class=card id=keywordCard style=display:none><h2>⚙️ 关键词管理</h2><p style=color:#94a3b8;font-size:.82rem;margin-bottom:8px>配置招标扫描关键词（当前仅扫描前 5 个 CCGP + 前 2 个 okcis）</p><div id=kwEdit></div><button class="btn btn-primary" onclick=saveKeywords() style=margin-top:8px>保存</button></div>
<div class=card><h2>📜 日志</h2><div id=logArea>暂无</div></div>
<script>
const kw=["水处理药剂","阻垢剂","杀菌剂","循环水处理","电厂药剂","反渗透","缓蚀剂","絮凝剂","脱盐水"];
document.getElementById('keywordArea').innerHTML=kw.map(k=>'<span class=keyword-tag>'+k+'</span>').join('');
async function load(){
const r=await fetch('/api/monitor/stats');const d=await r.json();
document.getElementById('statsArea').innerHTML='<div class=stat><div class=num>'+d.total+'</div><div class=label>总线索</div></div><div class=stat><div class=num>'+d.new+'</div><div class=label>新</div></div><div class=stat><div class=num>'+d.contacted+'</div><div class=label>已联系</div></div><div class=stat><div class=num>'+d.logs.length+'</div><div class=label>扫描</div></div>';
document.getElementById('itemList').innerHTML=d.recent.length?d.recent.map(i=>'<div class=item><div class=t><div class=title>'+i.title+'</div><div class=meta>'+(i.buyer?'🏢 '+i.buyer+' · ':'')+i.source+' '+new Date(i.time*1000).toLocaleString()+(i.contact?' · 👤 '+i.contact:'')+(i.email?' · 📧 '+i.email:'')+'</div>'+(i.phone?'<div style=color:#22c55e;font-size:.85rem>📞 '+i.phone+'</div>':'')+'</div><div style=display:flex;gap:4px;align-items:center>'+(i.status=='new'?'<button class="btn-ghost" onclick="markContacted('+"'"+'${i.id}'+"'"+')">📞 已联系</button>':'<span class="tag tag-'+i.status+'">'+i.status+'</span>')+'</div></div>').join(''):'<p style=color:#64748b>暂无</p>';
document.getElementById('logArea').innerHTML=d.logs.length?d.logs.map(l=>'<div>'+l.source+' 找到'+l.found+' 新增'+l.new+'</div>').join(''):'暂无';
}
async function runScan(){
document.getElementById('scanBtn').disabled=true;document.getElementById('scanBtn').textContent='⏳ 扫描中...';
const r=await fetch('/api/monitor/scan',{method:'POST'});const d=await r.json();
document.getElementById('scanBtn').disabled=false;document.getElementById('scanBtn').textContent='🔄 扫描 (+'+d.total_new+')';
load();}
load();setInterval(load,15000);
async function showKeywords(){const c=document.getElementById('keywordCard');c.style.display=c.style.display=='none'?'block':'none';const r=await fetch('/api/keywords');const d=await r.json();document.getElementById('kwEdit').innerHTML='<div style=display:flex;flex-wrap:wrap;gap:6px>'+(d.all||d.ccgp||[]).map(k=>'<label style=display:flex;align-items:center;gap:4px;font-size:.82rem;background:#0f172a;padding:4px 10px;border-radius:6px><input type=checkbox checked value="'+k+'">'+k+'</label>').join('')+'</div>';}
async function saveKeywords(){const boxes=document.querySelectorAll('#kwEdit input:checked');const kw=Array.from(boxes).map(b=>b.value);await fetch('/api/keywords',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({keywords:kw})});alert('已保存，下次扫描生效');}
function showCompany(name){fetch('/api/company/'+encodeURIComponent(name)).then(r=>r.json()).then(data=>{const m=document.createElement('div');m.className='modal';m.style.display='flex';m.innerHTML='<div class=modal-box style=max-width:700px><button class=modal-close onclick=this.parentElement.parentElement.style.display=\"none\">&times;</button><h2>\\ud83c\\udfe2 '+name+'</h2><p style=color:#94a3b8;margin-bottom:12px>\\u5171 '+data.length+'\\u6761\\u8bb0\\u5f55</p>'+data.map(i=>'<div class=item><div class=t><div class=title>'+i.title+'</div><div class=meta>'+(i.contact?'\\ud83d\\udc64'+i.contact:'')+' '+(i.phone?'\\ud83d\\udcde'+i.phone:'')+' '+(i.contact||i.phone?String.fromCharCode(183):'')+new Date(i.time*1000).toLocaleString()+'</div></div></div>').join('')+'</div>';m.onclick=e=>{if(e.target===m)m.remove()};document.body.appendChild(m)});}
</script>
</body></html>"""


@app.route("/")
def index():
    return render_template_string(INDEX_HTML)


@app.route("/monitor")
def monitor_page():
    return render_template_string(MONITOR_HTML)

@app.route("/api/monitor/stats")
def monitor_stats():
    return jsonify(get_stats())

@app.route("/api/monitor/scan", methods=["POST"])
def monitor_scan():
    result = scan_with_notify()
    return jsonify(result)

@app.route("/api/monitor/status", methods=["POST"])
def monitor_update_status():
    body = request.get_json(silent=True) or {}
    item_id = body.get("id")
    status = body.get("status", "new")
    notes = body.get("notes", "")
    if item_id:
        update_status(item_id, status, notes)
    return jsonify(success=True)


@app.route("/upload", methods=["POST"])
def upload():
    file = request.files.get("file")
    if not file:
        return jsonify(success=False, error="请选择文件")
    ext = Path(file.filename).suffix.lower()
    if ext not in (".xlsx", ".xls"):
        return jsonify(success=False, error="仅支持 .xlsx / .xls")
    tid = uuid.uuid4().hex[:8]
    path = UPLOAD_DIR / f"{tid}{ext}"
    file.save(str(path))
    try:
        wb = load_workbook(str(path), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if len(rows) < 2:
            return jsonify(success=False, error="文件为空")
        header = [str(c or "").strip() for c in rows[0]]
        name_col = None
        for i, h in enumerate(header):
            if "公司" in h or "企业" in h or "名称" in h or "name" in h.lower():
                name_col = i; break
        if name_col is None:
            return jsonify(success=False, error="未找到「公司名称」列")
        companies = []
        for row in rows[1:]:
            name = str(row[name_col] or "").strip()
            if name and len(name) > 1:
                companies.append({"name": name, "row": [str(c or "") for c in row], "header": header})
        if not companies:
            return jsonify(success=False, error="未找到有效公司数据")
        tasks[tid] = {"companies": companies, "results": [], "done": 0, "total": len(companies), "phones": 0, "failed": 0, "logs": [], "running": False}
        preview = [c["row"] for c in companies[:5]]
        return jsonify(success=True, task_id=tid, count=len(companies), columns=header, rows=preview)
    except Exception as e:
        return jsonify(success=False, error=f"解析失败: {str(e)}")

@app.route("/start/<tid>", methods=["POST"])
def start_search(tid):
    t = tasks.get(tid)
    if not t:
        return jsonify(success=False)
    t["running"] = True
    t["results"] = []
    t["done"] = 0
    t["phones"] = 0
    t["failed"] = 0
    t["logs"] = []
    def _run():
        for i, c in enumerate(t["companies"]):
            name = c["name"]
            t["logs"].append({"i": f"i{i}", "t": "info", "m": f"[{i+1}/{t['total']}] {name}"})
            try:
                result = search_company_phones(name)
                phones = result.get("phones", [])
                t["results"].append({"company": name, "phones": phones, "row": c["row"]})
                if phones:
                    t["phones"] += len(phones)
                    t["logs"].append({"i": f"p{i}", "t": "found", "m": f"  ✓ {len(phones)} 个: {' | '.join(phones[:3])}"})
                else:
                    t["failed"] += 1
                    t["logs"].append({"i": f"e{i}", "t": "empty", "m": f"  ✗ 未找到电话"})
            except Exception as e:
                t["failed"] += 1
                t["results"].append({"company": name, "phones": [], "row": c["row"]})
                t["logs"].append({"i": f"x{i}", "t": "empty", "m": f"  ✗ 错误: {str(e)[:50]}"})
            t["done"] += 1
            _save_result(t, tid)
        t["running"] = False
    threading.Thread(target=_run, daemon=True).start()
    return jsonify(success=True)

@app.route("/progress/<tid>")
def progress(tid):
    t = tasks.get(tid)
    if not t:
        return jsonify(done=0, total=0, percent=0, phones=0, failed=0, logs=[])
    percent = (t["done"] / t["total"] * 100) if t["total"] > 0 else 0
    return jsonify(done=t["done"], total=t["total"], percent=round(percent,1), phones=t["phones"], failed=t["failed"], logs=t["logs"][-50:])

@app.route("/download/<tid>")
def download(tid):
    path = UPLOAD_DIR / f"{tid}_result.xlsx"
    if path.exists():
        return send_file(str(path), as_attachment=True, download_name="获客结果.xlsx")
    return "文件未就绪", 404

def _save_result(t, tid):
    wb = Workbook()
    ws = wb.active
    ws.title = "获客结果"
    header = t["companies"][0]["header"] if t["companies"] else []
    ws.append(header + ["搜到电话", "电话数量"])
    for r in t["results"]:
        ws.append(r["row"] + [" | ".join(r["phones"]), len(r["phones"])])
    phone_col = len(header) + 1
    if phone_col:
        ws.cell(row=1, column=phone_col).value = "搜到电话"
        ws.column_dimensions[chr(64 + phone_col) if phone_col <= 26 else "A"].width = 40
    wb.save(str(UPLOAD_DIR / f"{tid}_result.xlsx"))


# ============ 招标导出 ============

@app.route("/api/monitor/export")
def monitor_export():
    """导出招标线索为 Excel"""
    stats = get_stats()
    wb = Workbook()
    ws = wb.active
    ws.title = "招标线索"
    ws.append(["标题", "采购单位", "联系人", "电话", "邮箱", "来源", "时间", "状态"])
    for item in stats["recent"]:
        ws.append([
            item.get("title", ""),
            item.get("buyer", ""),
            item.get("contact", ""),
            item.get("phone", ""),
            item.get("email", ""),
            item.get("source", ""),
            datetime.fromtimestamp(item.get("time", 0)).strftime("%Y-%m-%d %H:%M"),
            item.get("status", ""),
        ])
    ws.column_dimensions['A'].width = 50
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['D'].width = 25
    out = UPLOAD_DIR / "招标线索.xlsx"
    wb.save(str(out))
    return send_file(str(out), as_attachment=True, download_name="招标线索.xlsx")


# ============ 定时扫描 ============

import threading as _th
import time as _time

_SCAN_INTERVAL = 3600  # 1 小时
_last_scan_time = [0.0]

def _scheduler():
    """后台定时扫描招标"""
    _last_scan_time[0] = _time.time()
    while True:
        try:
            if _time.time() - _last_scan_time[0] >= _SCAN_INTERVAL:
                logger.info("[定时扫描] 开始...")
                from bidding_monitor import scan_with_notify, close_db
                r = scan_with_notify()
                close_db()
                logger.info(f"[定时扫描] 完成: {r['total_found']} 条, 新增 {r['new_saved']}")
                _last_scan_time[0] = _time.time()
        except Exception as e:
            logger.error(f"[定时扫描] 失败: {e}")
        _time.sleep(300)  # 每 5 分钟检查一次

# 启动后台线程
_th.Thread(target=_scheduler, daemon=True).start()
if __name__ == "__main__":
    print(f"BDSS 获客系统 http://127.0.0.1:5300")
    logger.info("定时扫描已启动 (间隔: 1 小时)")
    app.run(host="127.0.0.1", port=5300, debug=False)
    print(f"BDSS 获客系统 http://127.0.0.1:5300")
    app.run(host="127.0.0.1", port=5300, debug=False)
